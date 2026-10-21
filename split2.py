import pandas as pd
import os
import re
import shutil
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from collections import defaultdict
import warnings

warnings.filterwarnings("ignore")


# Load districts from CSV
def load_districts(csv_path="districts.csv"):
    """Load districts from CSV file"""
    districts_df = pd.read_csv(csv_path)
    district_names = districts_df["name"].str.upper().tolist()
    return districts_df, district_names


def create_phone_district_mapping():
    """
    Create a mapping of phone prefixes to districts
    """
    phone_mapping = {
        # Add actual mappings based on your data
    }
    return phone_mapping


def extract_district_from_text(text, district_names):
    """
    Extract district name from text by checking against district names
    """
    if pd.isna(text) or text == "":
        return None

    text_upper = str(text).upper()

    # Check for exact district name matches (longest first to avoid partial matches)
    sorted_districts = sorted(district_names, key=len, reverse=True)

    for district in sorted_districts:
        # Check for district name with word boundaries
        if re.search(r"\b" + re.escape(district) + r"\b", text_upper):
            return district

        # Check variations
        district_variations = [
            district.replace(" CITY", ""),
            district.replace(" ", ""),
            district.replace(" ", "-"),
            district.replace("-", " "),
        ]

        for variation in district_variations:
            if len(variation) > 3 and re.search(r"\b" + re.escape(variation) + r"\b", text_upper):
                return district

    return None


def find_phone_column(df):
    """
    Find the phone column in the dataframe
    """
    phone_col_names = [
        "phone",
        "phone_number",
        "phonenumber",
        "mobile",
        "contact",
        "telephone",
        "tel",
        "msisdn",
        "phone_no",
        "phone no",
        "cell",
    ]

    for col in df.columns:
        if col.lower().strip() in phone_col_names:
            return col

    # Check for columns containing 'phone' or 'mobile'
    for col in df.columns:
        if "phone" in col.lower() or "mobile" in col.lower() or "msisdn" in col.lower():
            return col

    return None


def normalize_phone(phone):
    """
    Normalize phone number to a standard format
    """
    if pd.isna(phone):
        return None

    phone_str = str(phone).strip()
    # Remove any non-digit characters
    phone_str = re.sub(r"\D", "", phone_str)

    # If it starts with 256, convert to 0
    if phone_str.startswith("256") and len(phone_str) > 9:
        phone_str = "0" + phone_str[3:]

    # If it's 9 digits, add leading 0
    if len(phone_str) == 9:
        phone_str = "0" + phone_str

    return phone_str


def get_district_for_payment(
    row, district_names, phone_mapping, phone_district_history=None, phone_col=None
):
    """
    Determine district for a payment based on narration, phone number, or other fields
    """
    # First priority: Check narration/text fields for district names
    text_fields = [
        "narration",
        "description",
        "remarks",
        "transaction_details",
        "notes",
        "comment",
        "details",
        "narrative",
        "particulars",
        "purpose",
    ]

    for field in text_fields:
        for col in row.index:
            if field.lower() in col.lower():
                if pd.notna(row[col]):
                    field_value = str(row[col])
                    district = extract_district_from_text(field_value, district_names)
                    if district:
                        return district

    # Second priority: Check all other text fields for district names
    for col in row.index:
        if col.lower() not in ["phone", "phone_number", "mobile", "contact", "district"]:
            if pd.notna(row[col]) and isinstance(row[col], str):
                district = extract_district_from_text(row[col], district_names)
                if district:
                    return district

    # Third priority: Check phone number against known phone-district history
    if phone_col and phone_col in row and pd.notna(row[phone_col]):
        phone_str = normalize_phone(row[phone_col])
        if phone_str and phone_district_history:
            for district, phones in phone_district_history.items():
                if phone_str in phones:
                    return district

    # Fourth priority: Check phone number prefix mapping
    if phone_col and phone_col in row and pd.notna(row[phone_col]):
        phone_str = normalize_phone(row[phone_col])
        if phone_str:
            for prefix, district in phone_mapping.items():
                if phone_str.startswith(prefix):
                    return district

    # Fifth priority: Check if district is already in data
    for col in row.index:
        if "district" in col.lower():
            if pd.notna(row[col]):
                district_name = str(row[col]).strip().upper()
                if district_name in district_names:
                    return district_name

    # If no district found, return as uncategorized
    return "UNCATEGORIZED"


def build_phone_district_history(all_dataframes, district_names, phone_col_name):
    """
    Build a comprehensive history of which phone numbers belong to which districts
    by scanning all data across all sheets
    """
    phone_district_history = defaultdict(set)

    if not phone_col_name:
        return phone_district_history

    # First pass: Find all phones associated with districts via narration
    for df in all_dataframes:
        if phone_col_name not in df.columns:
            continue

        for _, row in df.iterrows():
            if pd.notna(row[phone_col_name]):
                phone = normalize_phone(row[phone_col_name])
                if not phone:
                    continue

                # Try to find district from narration
                district = None
                text_fields = [
                    "narration",
                    "description",
                    "remarks",
                    "transaction_details",
                    "notes",
                    "comment",
                    "details",
                    "narrative",
                    "particulars",
                    "purpose",
                ]

                for field in text_fields:
                    for col in row.index:
                        if field.lower() in col.lower() and pd.notna(row[col]):
                            district = extract_district_from_text(str(row[col]), district_names)
                            if district:
                                break
                    if district:
                        break

                # If no district found in text fields, check all string columns
                if not district:
                    for col in row.index:
                        if pd.notna(row[col]) and isinstance(row[col], str):
                            district = extract_district_from_text(row[col], district_names)
                            if district:
                                break

                if district:
                    phone_district_history[district].add(phone)

    print(f"\n  Built phone-district history:")
    for district, phones in sorted(
        phone_district_history.items(), key=lambda x: len(x[1]), reverse=True
    ):
        print(f"    {district}: {len(phones)} unique phone numbers")

    return phone_district_history


def load_existing_district_data(district_folder, base_name, district):
    """
    Load existing data for a district if it exists
    """
    if not os.path.exists(district_folder):
        return None

    main_file = os.path.join(district_folder, f"{base_name}_{district}.xlsx")

    if os.path.exists(main_file):
        try:
            df = pd.read_excel(main_file)
            print(f"  Found existing file with {len(df)} records")
            return df
        except Exception as e:
            print(f"  Warning: Could not read existing file: {e}")
            return None

    part_files = [
        f
        for f in os.listdir(district_folder)
        if f.startswith(f"{base_name}_{district}_part_") and f.endswith(".xlsx")
    ]

    if part_files:
        print(f"  Found {len(part_files)} part files")
        combined_df = pd.DataFrame()
        for part_file in sorted(part_files):
            try:
                df = pd.read_excel(os.path.join(district_folder, part_file))
                combined_df = pd.concat([combined_df, df], ignore_index=True)
            except Exception as e:
                print(f"  Warning: Could not read {part_file}: {e}")

        if not combined_df.empty:
            print(f"  Combined data has {len(combined_df)} records")
            return combined_df

    return None


def append_to_existing_file(existing_df, new_data, district_folder, base_name, district):
    """
    Append new data to existing file and save
    """
    combined_data = pd.concat([existing_df, new_data], ignore_index=True)
    print(
        f"  Combined data: {len(combined_data)} records ({len(existing_df)} existing + {len(new_data)} new)"
    )

    filename = f"{base_name}_{district}.xlsx"
    filepath = os.path.join(district_folder, filename)

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        combined_data.to_excel(writer, sheet_name="Payments", index=False)

    # Adjust column widths
    try:
        wb = load_workbook(filepath)
        ws = wb.active

        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        wb.save(filepath)
    except Exception as e:
        print(f"  Warning: Could not adjust column widths: {e}")

    file_size = os.path.getsize(filepath) / (1024 * 1024)
    print(f"  Saved: {filename} ({file_size:.2f} MB)")

    if file_size > 1.9:
        print(f"  File is {file_size:.2f} MB, splitting into chunks...")
        split_file_into_chunks(combined_data, district_folder, base_name, district)

    return combined_data


def split_file_into_chunks(data, district_folder, base_name, district, max_size_mb=1.9):
    """
    Split data into chunks and save as separate files
    """
    for file in os.listdir(district_folder):
        if file.startswith(f"{base_name}_{district}_part_") and file.endswith(".xlsx"):
            os.remove(os.path.join(district_folder, file))

    chunks = split_into_chunks(data, max_size_mb)

    for i, chunk in enumerate(chunks, 1):
        chunk_filename = f"{base_name}_{district}_part_{i}.xlsx"
        chunk_filepath = os.path.join(district_folder, chunk_filename)

        with pd.ExcelWriter(chunk_filepath, engine="openpyxl") as writer:
            chunk.to_excel(writer, sheet_name="Payments", index=False)

        chunk_size = os.path.getsize(chunk_filepath) / (1024 * 1024)
        print(f"    Saved: {chunk_filename} ({chunk_size:.2f} MB)")


def split_into_chunks(df, max_size_mb=1.9):
    """
    Split dataframe into chunks based on approximate file size
    """
    if len(df) == 0:
        return []

    estimated_size_per_row = df.memory_usage(deep=True).sum() / len(df) / (1024 * 1024)
    rows_per_chunk = int((max_size_mb * 0.8) / estimated_size_per_row)
    rows_per_chunk = max(1, rows_per_chunk)

    chunks = []
    for i in range(0, len(df), rows_per_chunk):
        chunk = df.iloc[i : i + rows_per_chunk]
        chunks.append(chunk)

    return chunks


def create_backup(existing_file_path):
    """
    Create a backup of existing file before appending
    """
    if not os.path.exists(existing_file_path):
        return

    backup_folder = os.path.join(os.path.dirname(existing_file_path), "backup")
    os.makedirs(backup_folder, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = os.path.basename(existing_file_path)
    backup_name = f"{os.path.splitext(filename)[0]}_{timestamp}.xlsx"
    backup_path = os.path.join(backup_folder, backup_name)

    shutil.copy2(existing_file_path, backup_path)
    print(f"  Backup created: {backup_name}")


def process_all_sheets(excel_file, districts_csv="districts.csv"):
    """
    Process all sheets in an Excel file - handles multiple sheets
    """
    # Load districts
    districts_df, district_names = load_districts(districts_csv)

    if not os.path.exists(excel_file):
        print(f"File not found: {excel_file}")
        return None

    print(f"\n{'='*60}")
    print(f"Processing: {os.path.basename(excel_file)}")
    print(f"{'='*60}")

    # Read all sheets
    try:
        xl = pd.ExcelFile(excel_file)
        sheet_names = xl.sheet_names
        print(f"Sheets found: {sheet_names}")
    except Exception as e:
        print(f"Error reading {excel_file}: {e}")
        return None

    # Load all sheets into dataframes
    all_dataframes = {}
    phone_col_name = None

    for sheet_name in sheet_names:
        try:
            df = pd.read_excel(excel_file, sheet_name=sheet_name)
            if len(df) > 0:
                all_dataframes[sheet_name] = df
                print(
                    f"  Sheet '{sheet_name}': {len(df)} rows, columns: {', '.join(df.columns.tolist())}"
                )

                # Find phone column
                if not phone_col_name:
                    phone_col_name = find_phone_column(df)
                    if phone_col_name:
                        print(f"  Phone column detected: '{phone_col_name}'")
            else:
                print(f"  Sheet '{sheet_name}': empty, skipping")
        except Exception as e:
            print(f"  Error reading sheet '{sheet_name}': {e}")

    if not all_dataframes:
        print("No data found in any sheet")
        return None

    # Combine all sheets for phone history building
    all_dfs_list = list(all_dataframes.values())

    # Build phone-district history from all sheets combined
    print(f"\nBuilding phone-district history...")
    phone_mapping = create_phone_district_mapping()
    phone_district_history = build_phone_district_history(
        all_dfs_list, district_names, phone_col_name
    )

    # Create splits directory
    splits_dir = "C:/xampp/htdocs/ecashverifier/splits"
    os.makedirs(splits_dir, exist_ok=True)

    # Get base filename without path and extension
    base_name = os.path.splitext(os.path.basename(excel_file))[0]

    # Process each sheet
    all_district_data = defaultdict(list)  # district -> list of dataframes from different sheets

    for sheet_name, df in all_dataframes.items():
        print(f"\n{'─'*60}")
        print(f"Processing sheet: {sheet_name}")
        print(f"{'─'*60}")

        # Add district column
        def apply_district(row):
            return get_district_for_payment(
                row, district_names, phone_mapping, phone_district_history, phone_col_name
            )

        df = df.copy()
        df["district"] = df.apply(apply_district, axis=1)

        # Show distribution
        district_counts = df["district"].value_counts()
        print(f"  District distribution:")
        for district, count in district_counts.items():
            print(f"    {district}: {count}")

        # Group by district and collect
        for district in df["district"].unique():
            district_data = df[df["district"] == district].copy()
            district_data = district_data.drop("district", axis=1)
            all_district_data[district].append(district_data)

    # Now process each district by combining data from all sheets
    print(f"\n{'='*60}")
    print("Saving district files...")
    print(f"{'='*60}")

    summary_info = []
    total_records = 0

    for district, dataframes_list in all_district_data.items():
        # Combine data from all sheets for this district
        combined_new_data = pd.concat(dataframes_list, ignore_index=True)

        print(f"\nDistrict: {district} - {len(combined_new_data)} new records")

        if len(combined_new_data) == 0:
            continue

        # Create district folder
        district_folder = os.path.join(splits_dir, district)
        os.makedirs(district_folder, exist_ok=True)

        # Check if existing file exists
        main_file = os.path.join(district_folder, f"{base_name}_{district}.xlsx")
        existing_data = load_existing_district_data(district_folder, base_name, district)

        if existing_data is not None:
            create_backup(main_file)
            combined_data = append_to_existing_file(
                existing_data, combined_new_data, district_folder, base_name, district
            )
            record_count = len(combined_data)
            file_size = (
                os.path.getsize(main_file) / (1024 * 1024) if os.path.exists(main_file) else 0
            )
        else:
            print(f"  No existing data found, creating new file")
            filename = f"{base_name}_{district}.xlsx"
            filepath = os.path.join(district_folder, filename)

            with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
                combined_new_data.to_excel(writer, sheet_name="Payments", index=False)

            try:
                wb = load_workbook(filepath)
                ws = wb.active

                for column in ws.columns:
                    max_length = 0
                    column_letter = get_column_letter(column[0].column)
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    ws.column_dimensions[column_letter].width = adjusted_width

                wb.save(filepath)
            except Exception as e:
                print(f"  Warning: Could not adjust column widths: {e}")

            file_size = os.path.getsize(filepath) / (1024 * 1024)
            print(f"  Saved: {filename} ({file_size:.2f} MB)")

            if file_size > 1.9:
                print(f"  File is {file_size:.2f} MB, splitting into chunks...")
                split_file_into_chunks(combined_new_data, district_folder, base_name, district)

            record_count = len(combined_new_data)

        summary_info.append(
            {
                "District": district,
                "New Records": len(combined_new_data),
                "Total Records": record_count,
                "File Size (MB)": round(file_size, 2),
            }
        )
        total_records += record_count

    # Save summary report
    if summary_info:
        summary_df = pd.DataFrame(summary_info)
        summary_df = summary_df.sort_values("Total Records", ascending=False)

        summary_file = os.path.join(splits_dir, f"{base_name}_summary_report.xlsx")
        with pd.ExcelWriter(summary_file, engine="openpyxl") as writer:
            summary_df.to_excel(writer, sheet_name="Summary", index=False)

            stats_df = pd.DataFrame(
                {
                    "Statistic": [
                        "Total Records",
                        "Total Districts",
                        "New Records Added",
                        "Sheets Processed",
                        "Update Date",
                    ],
                    "Value": [
                        total_records,
                        len(all_district_data),
                        sum(len(df) for df in all_dataframes.values()),
                        len(all_dataframes),
                        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    ],
                }
            )
            stats_df.to_excel(writer, sheet_name="Statistics", index=False)

        print(f"\n{'='*60}")
        print(f"Summary report saved: {summary_file}")
        print(f"Total records across all districts: {total_records}")
        print(f"Sheets processed: {len(all_dataframes)}")
        print(f"\nDistrict breakdown:")
        print(summary_df.to_string(index=False))

    return all_district_data


# Main execution
if __name__ == "__main__":
    data_dir = "C:/xampp/htdocs/ecashverifier"

    print(f"Looking for files in: {data_dir}")
    print(f"Directory exists: {os.path.exists(data_dir)}")

    if os.path.exists(data_dir):
        all_files = os.listdir(data_dir)
        print(f"\nAll files in directory:")
        for f in all_files:
            print(f"  - {f}")

    # Find all Excel files
    excel_files = []
    if os.path.exists(data_dir):
        for f in os.listdir(data_dir):
            if f.endswith(".xlsx") or f.endswith(".xls"):
                if f.startswith("~$"):
                    continue
                if f.startswith("splits"):
                    continue
                if "backup" in f.lower():
                    continue
                if "summary_report" in f.lower():
                    continue
                if "_part_" in f:
                    continue
                excel_files.append(f)

    if not excel_files:
        print("\nNo Excel files found in C:/xampp/htdocs/ecashverifier")
        exit()

    print(f"\nFound Excel files to process: {excel_files}")
    print("\nProcessing files (all sheets will be processed)...")
    print("=" * 60)

    for excel_file in excel_files:
        file_path = os.path.join(data_dir, excel_file)
        try:
            process_all_sheets(file_path, "districts.csv")
        except Exception as e:
            print(f"Error processing {excel_file}: {e}")
            import traceback

            traceback.print_exc()

    print(f"\n{'='*60}")
    print("ALL FILES PROCESSED SUCCESSFULLY")
    print(f"{'='*60}")
    print(f"All output files are in: {data_dir}/splits/")
